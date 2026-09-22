"""Build the ready-to-upload Google Sheets workbook.

Produces device-trade-in-pricing.xlsx with three tabs:

    Devices     -> one row per (device_id, condition), live price formulas
    Conditions  -> the four conditions and their multipliers
    Guide       -> legend, price math, and provenance of the sample prices

Upload the result to Google Drive and open it with Google Sheets; the tabs,
formulas, formatting and dropdowns all carry over.

Device and condition data is imported from generate_sample_data.py so there is
one source of truth for the sample catalogue.

Run:  python tools/build_workbook.py
"""

import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generate_sample_data import CONDITIONS, DEVICE_HEADER, DEVICES, LAST_UPDATED

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT = os.path.join(ROOT, "device-trade-in-pricing.xlsx")

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
INPUT_FONT = Font(name=FONT, color="0000FF")      # hand-edited values
FORMULA_FONT = Font(name=FONT, color="000000")    # calculated, do not overtype
BODY_FONT = Font(name=FONT)
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(bottom=THIN)

EUR = '€#,##0'

COLUMN_WIDTHS = {
    "A": 16, "B": 14, "C": 22, "D": 10, "E": 13, "F": 16,
    "G": 12, "H": 16, "I": 15, "J": 9, "K": 14,
}


def style_header(sheet, width_map):
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    for column, width in width_map.items():
        sheet.column_dimensions[column].width = width


def build_devices(sheet):
    sheet.append(DEVICE_HEADER)

    row_number = 1
    for device_id, manufacturer, model, storage, year, base in DEVICES:
        for condition, multiplier, _desc in CONDITIONS:
            row_number += 1
            sheet.append([
                device_id, manufacturer, model, storage, year,
                base, condition, multiplier,
                # Column F = base_price_eur, column H = price_multiplier.
                f"=ROUND(F{row_number}*H{row_number},0)",
                "TRUE", LAST_UPDATED,
            ])

    last_row = row_number
    for row in sheet.iter_rows(min_row=2, max_row=last_row, max_col=11):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
        row[4].alignment = Alignment(horizontal="center")   # release_year
        row[5].font = INPUT_FONT                            # base_price_eur
        row[5].number_format = EUR
        row[7].number_format = "0.00"                       # price_multiplier
        row[8].font = FORMULA_FONT                          # final_price_eur
        row[8].number_format = EUR
        row[9].alignment = Alignment(horizontal="center")   # active

    style_header(sheet, COLUMN_WIDTHS)
    sheet.auto_filter.ref = f"A1:K{last_row}"

    # Dropdowns, extended past the sample data so added rows inherit them.
    condition_rule = DataValidation(
        type="list", formula1=f"Conditions!$A$2:$A${len(CONDITIONS) + 1}",
        allow_blank=False, showDropDown=False,
        errorTitle="Unknown condition",
        error="Pick one of the four conditions on the Conditions tab.",
    )
    sheet.add_data_validation(condition_rule)
    condition_rule.add(f"G2:G{last_row + 500}")

    active_rule = DataValidation(
        type="list", formula1='"TRUE,FALSE"', allow_blank=False,
        showDropDown=False, errorTitle="Invalid value",
        error="active must be TRUE or FALSE.",
    )
    sheet.add_data_validation(active_rule)
    active_rule.add(f"J2:J{last_row + 500}")

    return last_row


def build_conditions(sheet):
    sheet.append(["condition", "price_multiplier", "description"])
    for condition, multiplier, description in CONDITIONS:
        sheet.append([condition, multiplier, description])

    for row in sheet.iter_rows(min_row=2, max_row=len(CONDITIONS) + 1, max_col=3):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
        row[1].font = INPUT_FONT
        row[1].number_format = "0.00"
        row[2].alignment = Alignment(wrap_text=True, vertical="center")

    style_header(sheet, {"A": 14, "B": 17, "C": 66})


GUIDE_ROWS = [
    ("Device Trade-In Pricing", "heading"),
    ("", None),
    ("How the price is calculated", "subheading"),
    ("final_price_eur  =  ROUND(base_price_eur × price_multiplier, 0)", None),
    ("The Devices tab holds one row per device and condition, so every lookup of "
     "manufacturer + model + storage + condition lands on exactly one row.", None),
    ("", None),
    ("Worked example — iPhone 15 Pro 128GB, base price €520", "subheading"),
    ("Like New  × 1.00  =  €520", None),
    ("Intact    × 0.85  =  €442", None),
    ("Cracked   × 0.60  =  €312", None),
    ("Faulty    × 0.20  =  €104", None),
    ("", None),
    ("Which cells to edit", "subheading"),
    ("Blue text  — type here. base_price_eur (column F) is the only price you set by hand.", "input"),
    ("Black text — calculated. final_price_eur (column I) is a formula; typing a number "
     "over it breaks the link to the base price.", "formula"),
    ("Set active (column J) to FALSE to hide a device from the chatbot without deleting "
     "its pricing history.", None),
    ("", None),
    ("Adding a device", "subheading"),
    ("Add four rows, one per condition. Fastest way: select the four rows of a similar "
     "device, copy, paste at the bottom, then edit columns A to F. The formula in column I "
     "comes along and renumbers itself.", None),
    ("", None),
    ("Where these prices come from", "subheading"),
    ("The base prices in this workbook are realistic illustrative estimates for structure "
     "and testing. They are not sourced from KSP and should be replaced with real trade-in "
     "values before the chatbot quotes them to customers.", None),
    ("The multipliers (1.00 / 0.85 / 0.60 / 0.20) were specified for this project; they are "
     "stored per row rather than looked up live, so a past quote stays reproducible if the "
     "multipliers are retuned later.", None),
]


def build_guide(sheet):
    styles = {
        "heading": Font(name=FONT, bold=True, size=16, color="1F3864"),
        "subheading": Font(name=FONT, bold=True, size=11, color="1F3864"),
        "input": Font(name=FONT, color="0000FF"),
        "formula": Font(name=FONT, color="000000"),
    }
    for index, (text, style) in enumerate(GUIDE_ROWS, start=1):
        cell = sheet.cell(row=index, column=1, value=text)
        cell.font = styles.get(style, BODY_FONT)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style in (None, "input", "formula") and len(text) > 90:
            sheet.row_dimensions[index].height = 30
    sheet.column_dimensions["A"].width = 110
    sheet.sheet_view.showGridLines = False


if __name__ == "__main__":
    workbook = Workbook()

    devices = workbook.active
    devices.title = "Devices"
    row_count = build_devices(devices)

    build_conditions(workbook.create_sheet("Conditions"))
    build_guide(workbook.create_sheet("Guide"))

    workbook.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({row_count - 1} device rows)")
